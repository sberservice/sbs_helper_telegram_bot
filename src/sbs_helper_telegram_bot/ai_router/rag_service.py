"""
rag_service.py — сервис базы знаний (RAG) для AI-маршрутизации.

Отвечает за:
- приём и парсинг документов,
- разбиение на чанки,
- хранение чанков в БД,
- retrieval релевантных фрагментов,
- кэширование ответов.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
import asyncio
import math
import sys
import time
from collections import Counter
from dataclasses import dataclass
from typing import Any, Awaitable, Callable, Dict, List, Optional, Tuple, TypeVar

import src.common.database as database

from src.sbs_helper_telegram_bot.ai_router import settings as ai_settings
from src.sbs_helper_telegram_bot.ai_router.messages import (
    AI_PROGRESS_STAGE_RAG_AUGMENTED_REQUEST_STARTED,
    AI_PROGRESS_STAGE_RAG_PREFILTER_STARTED,
)
from src.sbs_helper_telegram_bot.ai_router.prompts import build_rag_prompt, build_rag_summary_prompt
from src.sbs_helper_telegram_bot.ai_router.vector_search import (
    LocalEmbeddingProvider,
    LocalVectorIndex,
)

logger = logging.getLogger(__name__)

try:
    from rank_bm25 import BM25Okapi
except Exception:
    BM25Okapi = None

_TOKEN_RE = re.compile(r"[a-zа-яё0-9]{3,}", re.IGNORECASE)
_CYRILLIC_TOKEN_RE = re.compile(r"[а-яё]", re.IGNORECASE)
_SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md", ".html", ".htm"}
_RAG_CHUNK_SCAN_LIMIT = 3000
_RAG_SUMMARY_SCAN_LIMIT = 3000
_RAG_EMBEDDING_UPSERT_BATCH_SIZE = 25
_RAG_EMBEDDING_UPSERT_MAX_RETRIES = 3
_RAG_EMBEDDING_RETRY_BASE_DELAY_SECONDS = 0.25
_MYSQL_RETRYABLE_ERRNOS = {1205, 1213}
_RAG_DB_OPERATION_MAX_RETRIES = 3
_RAG_DB_OPERATION_RETRY_BASE_DELAY_SECONDS = 0.25
_SPACES_RE = re.compile(r"\s+")

TResult = TypeVar("TResult")


@dataclass
class CachedAnswer:
    """Элемент TTL-кэша ответа RAG."""

    answer: str
    expires_at: float


class RagKnowledgeService:
    """Сервис работы с базой знаний RAG."""

    def __init__(self, cache_ttl_seconds: int = ai_settings.AI_RAG_CACHE_TTL_SECONDS):
        self._cache_ttl_seconds = cache_ttl_seconds
        self._answer_cache: Dict[str, CachedAnswer] = {}
        self._embedding_provider: Optional[LocalEmbeddingProvider] = None
        self._vector_index: Optional[LocalVectorIndex] = None
        self._summary_embedding_cache: Dict[int, List[float]] = {}
        self._summary_embedding_corpus_version: int = -1
        self._summary_vector_prefilter_source: str = "disabled"
        self._summary_vector_prefilter_hits: int = 0
        self._ru_morph_analyzer: Optional[object] = None
        self._ru_stemmer: Optional[object] = None
        self._normalized_token_cache: Dict[str, str] = {}
        self._normalization_dependency_warning_logged: bool = False

    @staticmethod
    def is_supported_file(filename: str) -> bool:
        """Проверить поддерживаемое расширение файла."""
        lower_name = (filename or "").lower()
        return any(lower_name.endswith(ext) for ext in _SUPPORTED_EXTENSIONS)

    @classmethod
    def _resolve_text_slicer_name(cls) -> str:
        """Определить имя активного slicer-а для текстового chunking."""
        if not cls._is_langchain_splitter_supported():
            return "manual_window_slicer"

        try:
            from langchain.text_splitter import RecursiveCharacterTextSplitter  # noqa: F401

            return "RecursiveCharacterTextSplitter(langchain.text_splitter)"
        except Exception:
            return "manual_window_slicer"

    def get_chunking_diagnostics(self) -> Dict[str, object]:
        """Вернуть текущую диагностику стратегии чанкинга для runtime-логирования."""
        html_splitter_enabled = ai_settings.is_rag_html_splitter_enabled()
        langchain_supported = self._is_langchain_splitter_supported()

        return {
            "chunk_size": int(ai_settings.AI_RAG_CHUNK_SIZE),
            "chunk_overlap": int(ai_settings.AI_RAG_CHUNK_OVERLAP),
            "html_splitter_enabled": bool(html_splitter_enabled),
            "langchain_splitter_supported": bool(langchain_supported),
            "text_slicer": self._resolve_text_slicer_name(),
            "html_strategy": (
                "html_semantic_preserving_splitter_with_fallback"
                if html_splitter_enabled
                else "plain_text_fallback(html_splitter_disabled)"
            ),
            "plain_text_strategy": "extract_text_then_split_text",
        }

    def _log_chunking_strategy(
        self,
        *,
        file_name: str,
        file_format: str,
        strategy: str,
        chunks_count: int,
    ) -> None:
        """Записать в лог выбранную стратегию chunking для документа."""
        diagnostics = self.get_chunking_diagnostics()
        logger.info(
            "RAG chunking strategy: file=%s format=%s strategy=%s slicer=%s chunk_size=%s chunk_overlap=%s chunks=%s html_splitter_enabled=%s langchain_splitter_supported=%s",
            file_name,
            file_format,
            strategy,
            diagnostics.get("text_slicer"),
            diagnostics.get("chunk_size"),
            diagnostics.get("chunk_overlap"),
            chunks_count,
            diagnostics.get("html_splitter_enabled"),
            diagnostics.get("langchain_splitter_supported"),
        )

    def _execute_with_db_retry(self, operation_name: str, operation: Callable[[], TResult]) -> TResult:
        """Выполнить DB-операцию с retry для временных ошибок блокировок MySQL."""
        for attempt in range(1, _RAG_DB_OPERATION_MAX_RETRIES + 1):
            try:
                return operation()
            except Exception as exc:
                mysql_errno = getattr(exc, "errno", None)
                is_retryable = mysql_errno in _MYSQL_RETRYABLE_ERRNOS
                if not is_retryable or attempt >= _RAG_DB_OPERATION_MAX_RETRIES:
                    raise

                delay = _RAG_DB_OPERATION_RETRY_BASE_DELAY_SECONDS * (2 ** (attempt - 1))
                logger.warning(
                    "Повтор DB-операции после временной ошибки MySQL: operation=%s errno=%s attempt=%s/%s sleep=%.2fs",
                    operation_name,
                    mysql_errno,
                    attempt,
                    _RAG_DB_OPERATION_MAX_RETRIES,
                    delay,
                )
                time.sleep(delay)

        raise RuntimeError(f"DB retry loop exhausted: {operation_name}")

    def ingest_document_from_bytes(
        self,
        filename: str,
        payload: bytes,
        uploaded_by: int,
        source_type: str = "telegram",
        source_url: Optional[str] = None,
        upsert_vectors: bool = True,
        summary_model_scope: str = "default",
    ) -> Dict[str, int]:
        """
        Загрузить документ в базу знаний.

        Args:
            filename: Имя файла.
            payload: Байтовое содержимое файла.
            uploaded_by: Telegram ID администратора.
            source_type: Тип источника.
            source_url: URL источника (если применимо).
            upsert_vectors: Выполнять ли немедленный upsert векторных эмбеддингов.
            summary_model_scope: Контекст выбора модели summary (например, directory_ingest).

        Returns:
            Статистика загрузки документа.
        """
        if not filename:
            raise ValueError("Пустое имя файла")

        if not self.is_supported_file(filename):
            raise ValueError("Неподдерживаемый формат файла")

        max_bytes = ai_settings.AI_RAG_MAX_FILE_SIZE_MB * 1024 * 1024
        if len(payload) > max_bytes:
            raise ValueError("Файл слишком большой")

        content_hash = hashlib.sha256(payload).hexdigest()

        _reactivated_document_id: Optional[int] = None

        def _find_or_reactivate_existing_document() -> Optional[Dict[str, int]]:
            nonlocal _reactivated_document_id
            _reactivated_document_id = None
            with database.get_db_connection() as conn:
                with database.get_cursor(conn) as cursor:
                    cursor.execute(
                        """
                        SELECT id, status FROM rag_documents
                        WHERE content_hash = %s
                        LIMIT 1
                        """,
                        (content_hash,),
                    )
                    existing = cursor.fetchone()
                    if not existing:
                        return None

                    existing_id = int(existing["id"])
                    existing_status = str(existing.get("status") or "")

                    if existing_status != "active":
                        cursor.execute(
                            """
                            UPDATE rag_documents
                            SET
                                filename = %s,
                                source_type = %s,
                                source_url = %s,
                                uploaded_by = %s,
                                status = 'active',
                                updated_at = NOW()
                            WHERE id = %s
                            """,
                            (filename, source_type, source_url, uploaded_by, existing_id),
                        )
                        self._bump_corpus_version(
                            cursor,
                            f"reactivate:{existing_id}:{uploaded_by}:{filename}",
                        )
                        _reactivated_document_id = existing_id
                        logger.info(
                            "RAG ingest reactivated existing document: file=%s document_id=%s old_status=%s uploaded_by=%s",
                            filename,
                            existing_id,
                            existing_status,
                            uploaded_by,
                        )
                        return {
                            "document_id": existing_id,
                            "chunks_count": 0,
                            "is_duplicate": 1,
                            "reactivated": 1,
                        }

                    return {
                        "document_id": existing_id,
                        "chunks_count": 0,
                        "is_duplicate": 1,
                        "reactivated": 0,
                    }

        existing_result = self._execute_with_db_retry(
            operation_name="ingest.find_or_reactivate_by_hash",
            operation=_find_or_reactivate_existing_document,
        )
        if existing_result:
            if _reactivated_document_id is not None:
                self._set_vector_document_status(_reactivated_document_id, "active")
                self._clear_expired_cache()
            return existing_result

        if self._is_html_file(filename):
            chunks = self._split_html_payload(payload, filename=filename)
        else:
            extracted_text = self._extract_text(filename, payload)
            chunks = self._split_text(extracted_text)
            self._log_chunking_strategy(
                file_name=filename,
                file_format="text",
                strategy="extract_text_then_split_text",
                chunks_count=len(chunks),
            )

        if not chunks:
            raise ValueError("В документе не найден полезный текст")

        limited_chunks = chunks[: ai_settings.AI_RAG_MAX_CHUNKS_PER_DOC]
        summary_text, summary_model_name = self._generate_document_summary(
            filename,
            limited_chunks,
            user_id=uploaded_by,
            summary_model_scope=summary_model_scope,
        )

        def _insert_document_and_chunks() -> Tuple[int, List[Dict[str, object]]]:
            with database.get_db_connection() as conn:
                with database.get_cursor(conn) as cursor:
                    cursor.execute(
                        """
                        INSERT INTO rag_documents
                            (filename, source_type, source_url, uploaded_by, status, content_hash, created_at, updated_at)
                        VALUES (%s, %s, %s, %s, 'active', %s, NOW(), NOW())
                        """,
                        (filename, source_type, source_url, uploaded_by, content_hash),
                    )
                    local_document_id = int(cursor.lastrowid)
                    local_inserted_vector_chunks: List[Dict[str, object]] = []

                    for idx, chunk in enumerate(limited_chunks):
                        cursor.execute(
                            """
                            INSERT INTO rag_chunks
                                (document_id, chunk_index, chunk_text, created_at)
                            VALUES (%s, %s, %s, NOW())
                            """,
                            (local_document_id, idx, chunk),
                        )
                        local_inserted_vector_chunks.append(
                            {
                                "document_id": local_document_id,
                                "chunk_index": idx,
                                "filename": filename,
                                "chunk_text": chunk,
                                "status": "active",
                            }
                        )

                    self._upsert_document_summary(
                        cursor=cursor,
                        document_id=local_document_id,
                        summary_text=summary_text,
                        model_name=summary_model_name,
                    )
                    self._bump_corpus_version(cursor, f"upload:{filename}")

            return local_document_id, local_inserted_vector_chunks

        document_id, inserted_vector_chunks = self._execute_with_db_retry(
            operation_name="ingest.insert_document_and_chunks",
            operation=_insert_document_and_chunks,
        )

        if upsert_vectors:
            self._upsert_vectors_for_chunks(inserted_vector_chunks)

        self._clear_expired_cache()
        logger.info(
            "RAG ingest success: file=%s document_id=%s chunks=%s uploaded_by=%s",
            filename,
            document_id,
            len(limited_chunks),
            uploaded_by,
        )

        return {
            "document_id": document_id,
            "chunks_count": len(limited_chunks),
            "is_duplicate": 0,
        }

    def _generate_document_summary(
        self,
        filename: str,
        chunks: List[str],
        user_id: Optional[int] = None,
        summary_model_scope: str = "default",
    ) -> Tuple[str, Optional[str]]:
        """Сгенерировать summary документа с fallback на extractive-режим."""
        fallback_summary = self._build_fallback_summary(chunks)
        if not ai_settings.AI_RAG_SUMMARY_ENABLED:
            return fallback_summary, None

        excerpt = self._build_summary_excerpt(chunks)
        if not excerpt:
            return fallback_summary, None

        try:
            running_loop = asyncio.get_running_loop()
            if running_loop and running_loop.is_running():
                return fallback_summary, None
        except RuntimeError:
            pass

        try:
            from src.sbs_helper_telegram_bot.ai_router.llm_provider import get_provider

            provider = get_provider()
            model_override: Optional[str] = None
            if summary_model_scope == "directory_ingest":
                model_override = ai_settings.get_directory_ingest_summary_model_override()
                if model_override is None and ai_settings.AI_RAG_DIRECTORY_INGEST_SUMMARY_MODEL:
                    logger.warning(
                        "Некорректное значение AI_RAG_DIRECTORY_INGEST_SUMMARY_MODEL=%s, используется модель ответов по умолчанию",
                        ai_settings.AI_RAG_DIRECTORY_INGEST_SUMMARY_MODEL,
                    )

            system_prompt = build_rag_summary_prompt(
                document_name=filename,
                document_excerpt=excerpt,
                max_summary_chars=ai_settings.AI_RAG_SUMMARY_MAX_CHARS,
            )

            async def _request_summary() -> str:
                return await provider.chat(
                    messages=[{"role": "user", "content": "Сформируй summary документа. Подчеркни объект на который фокусируется документ. Подумай, какие в нем могут быть отличия в сути от других подобных похожих документов, чтобы далее было проще найти именно этот документ."}],
                    system_prompt=system_prompt,
                    user_id=user_id,
                    purpose="rag_summary",
                    model_override=model_override,
                )

            raw_summary = asyncio.run(_request_summary())
            normalized_summary = self._normalize_summary_text(raw_summary)
            if normalized_summary:
                if model_override:
                    return normalized_summary, ai_settings.normalize_deepseek_model(model_override)
                return normalized_summary, provider.get_model_name(purpose="rag_summary")
        except Exception as exc:
            logger.warning("Не удалось сгенерировать AI-summary для %s: %s", filename, exc)

        return fallback_summary, None

    @staticmethod
    def _build_summary_excerpt(chunks: List[str]) -> str:
        """Собрать ограниченный по длине фрагмент документа для суммаризации."""
        if not chunks:
            return ""

        max_chars = max(500, int(ai_settings.AI_RAG_SUMMARY_INPUT_MAX_CHARS))
        collected: List[str] = []
        current_len = 0

        for chunk in chunks:
            normalized_chunk = (chunk or "").strip()
            if not normalized_chunk:
                continue

            remaining = max_chars - current_len
            if remaining <= 0:
                break

            piece = normalized_chunk[:remaining]
            if piece:
                collected.append(piece)
                current_len += len(piece)

        return "\n\n".join(collected).strip()

    @staticmethod
    def _build_fallback_summary(chunks: List[str]) -> str:
        """Сформировать детерминированный fallback-summary из первых предложений документа."""
        excerpt = RagKnowledgeService._build_summary_excerpt(chunks)
        if not excerpt:
            return "Краткое summary недоступно: в документе не найден информативный текст."

        sentence_parts = re.split(r"(?<=[.!?])\s+", excerpt)
        selected = [part.strip() for part in sentence_parts if part.strip()][:8]
        if not selected:
            selected = [excerpt[: ai_settings.AI_RAG_SUMMARY_MAX_CHARS].strip()]

        fallback_text = " ".join(selected).strip()
        max_chars = max(200, int(ai_settings.AI_RAG_SUMMARY_MAX_CHARS))
        return fallback_text[:max_chars].strip()

    @staticmethod
    def _normalize_summary_text(summary_text: str) -> str:
        """Нормализовать итоговый текст summary перед сохранением."""
        normalized = re.sub(r"\s+", " ", (summary_text or "").strip())
        if not normalized:
            return ""
        max_chars = max(200, int(ai_settings.AI_RAG_SUMMARY_MAX_CHARS))
        return normalized[:max_chars].strip()

    def _upsert_document_summary(self, cursor, document_id: int, summary_text: str, model_name: Optional[str]) -> None:
        """Создать или обновить summary документа в rag_document_summaries."""
        safe_summary = (summary_text or "").strip()
        if not safe_summary:
            return

        cursor.execute(
            """
            INSERT INTO rag_document_summaries (document_id, summary_text, model_name, created_at, updated_at)
            VALUES (%s, %s, %s, NOW(), NOW())
            ON DUPLICATE KEY UPDATE
                summary_text = VALUES(summary_text),
                model_name = VALUES(model_name),
                updated_at = NOW()
            """,
            (document_id, safe_summary, model_name),
        )
        self._mark_summary_embedding_stale(cursor=cursor, document_id=document_id)

    @staticmethod
    def _mark_summary_embedding_stale(cursor, document_id: int) -> None:
        """Пометить summary-эмбеддинг документа как устаревший после обновления summary."""
        if document_id <= 0:
            return
        try:
            cursor.execute(
                """
                UPDATE rag_summary_embeddings
                SET embedding_status = 'stale', updated_at = NOW()
                WHERE document_id = %s
                """,
                (document_id,),
            )
        except Exception as exc:
            logger.warning("Не удалось пометить stale rag_summary_embeddings для document_id=%s: %s", document_id, exc)

    def list_documents(self, status: Optional[str] = None, limit: int = 20) -> List[Dict[str, object]]:
        """
        Получить список документов базы знаний.

        Args:
            status: Фильтр по статусу (active/archived/deleted) или None.
            limit: Максимальное число документов.

        Returns:
            Список документов с агрегированной статистикой.
        """
        safe_limit = max(1, min(limit, 200))
        statuses = {"active", "archived", "deleted"}

        query = """
            SELECT
                d.id,
                d.filename,
                d.source_type,
                d.source_url,
                d.uploaded_by,
                d.status,
                d.created_at,
                d.updated_at,
                COUNT(c.id) AS chunks_count
            FROM rag_documents d
            LEFT JOIN rag_chunks c ON c.document_id = d.id
        """
        params: List[object] = []

        if status and status in statuses:
            query += " WHERE d.status = %s"
            params.append(status)

        query += " GROUP BY d.id ORDER BY d.id DESC LIMIT %s"
        params.append(safe_limit)

        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                cursor.execute(query, tuple(params))
                rows = cursor.fetchall() or []

        return rows

    def list_documents_by_source(
        self,
        source_type: str,
        source_url_prefix: Optional[str] = None,
    ) -> List[Dict[str, object]]:
        """
        Получить документы по типу источника и префиксу source_url.

        Args:
            source_type: Тип источника (например, filesystem).
            source_url_prefix: Префикс source_url для фильтрации.

        Returns:
            Список документов с базовыми метаданными.
        """
        normalized_source_type = (source_type or "").strip()
        if not normalized_source_type:
            raise ValueError("Пустой source_type")

        query = """
            SELECT
                id,
                filename,
                source_type,
                source_url,
                uploaded_by,
                status,
                content_hash,
                created_at,
                updated_at
            FROM rag_documents
            WHERE source_type = %s
        """
        params: List[object] = [normalized_source_type]

        normalized_prefix = (source_url_prefix or "").strip()
        if normalized_prefix:
            query += " AND source_url LIKE %s"
            params.append(f"{normalized_prefix}%")

        query += " ORDER BY id DESC"

        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                cursor.execute(query, tuple(params))
                return cursor.fetchall() or []

    def get_document(self, document_id: int) -> Optional[Dict[str, object]]:
        """
        Получить детальную информацию по документу.

        Args:
            document_id: Идентификатор документа.

        Returns:
            Словарь с полями документа или None.
        """
        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                cursor.execute(
                    """
                    SELECT
                        d.id,
                        d.filename,
                        d.source_type,
                        d.source_url,
                        d.uploaded_by,
                        d.status,
                        d.content_hash,
                        d.created_at,
                        d.updated_at,
                        COUNT(c.id) AS chunks_count
                    FROM rag_documents d
                    LEFT JOIN rag_chunks c ON c.document_id = d.id
                    WHERE d.id = %s
                    GROUP BY d.id
                    LIMIT 1
                    """,
                    (document_id,),
                )
                return cursor.fetchone()

    def set_document_status(self, document_id: int, new_status: str, updated_by: int) -> bool:
        """
        Изменить статус документа (архивация/восстановление/мягкое удаление).

        Args:
            document_id: Идентификатор документа.
            new_status: Новый статус (active/archived/deleted).
            updated_by: ID администратора, выполняющего действие.

        Returns:
            True при успешном изменении, False если документ не найден.
        """
        allowed_statuses = {"active", "archived", "deleted"}
        if new_status not in allowed_statuses:
            raise ValueError("Некорректный статус документа")

        def _update_status_in_db() -> Optional[Tuple[bool, str]]:
            """Обновить статус в MySQL. Возвращает (changed, old_status) или None."""
            with database.get_db_connection() as conn:
                with database.get_cursor(conn) as cursor:
                    cursor.execute(
                        "SELECT status, filename FROM rag_documents WHERE id = %s LIMIT 1",
                        (document_id,),
                    )
                    existing = cursor.fetchone()
                    if not existing:
                        return None

                    old_status = str(existing.get("status", ""))
                    filename = str(existing.get("filename", "document"))

                    if old_status == new_status:
                        return (False, old_status)

                    cursor.execute(
                        """
                        UPDATE rag_documents
                        SET status = %s, updated_at = NOW()
                        WHERE id = %s
                        """,
                        (new_status, document_id),
                    )

                    if old_status == "active" or new_status == "active":
                        self._bump_corpus_version(
                            cursor,
                            f"status:{document_id}:{old_status}->{new_status}:{updated_by}:{filename}",
                        )

            return (True, old_status)

        result = self._execute_with_db_retry("set_document_status", _update_status_in_db)
        if result is None:
            return False

        changed, old_status = result
        if not changed:
            return True

        self._set_vector_document_status(document_id, new_status)
        self._clear_expired_cache()
        logger.info(
            "RAG document status changed: document_id=%s old=%s new=%s by=%s",
            document_id,
            old_status,
            new_status,
            updated_by,
        )
        return True

    def delete_document(self, document_id: int, updated_by: int, hard_delete: bool = False) -> bool:
        """
        Удалить документ из базы знаний.

        Args:
            document_id: Идентификатор документа.
            updated_by: ID администратора.
            hard_delete: True для физического удаления, иначе soft-delete.

        Returns:
            True если действие выполнено, иначе False.
        """
        if not hard_delete:
            return self.set_document_status(document_id, "deleted", updated_by)

        def _delete_from_db() -> bool:
            """Удалить документ из MySQL. Возвращает True если удалён, False если не найден."""
            with database.get_db_connection() as conn:
                with database.get_cursor(conn) as cursor:
                    cursor.execute(
                        "SELECT status, filename FROM rag_documents WHERE id = %s LIMIT 1",
                        (document_id,),
                    )
                    existing = cursor.fetchone()
                    if not existing:
                        return False

                    old_status = str(existing.get("status", ""))
                    filename = str(existing.get("filename", "document"))

                    cursor.execute("DELETE FROM rag_documents WHERE id = %s", (document_id,))
                    if old_status == "active":
                        self._bump_corpus_version(
                            cursor,
                            f"hard_delete:{document_id}:{updated_by}:{filename}",
                        )
            return True

        deleted = self._execute_with_db_retry("delete_document", _delete_from_db)
        if not deleted:
            return False

        self._delete_vector_document(document_id)
        self._clear_expired_cache()
        logger.info(
            "RAG document deleted: document_id=%s hard_delete=%s by=%s",
            document_id,
            hard_delete,
            updated_by,
        )
        return True

    async def answer_question(
        self,
        question: str,
        user_id: int,
        on_progress: Optional[Callable[[str, Optional[Dict[str, Any]]], Awaitable[None]]] = None,
    ) -> Optional[str]:
        """
        Ответить на вопрос пользователя на основе документов.

        Args:
            question: Текст вопроса.
            user_id: Telegram ID пользователя.

        Returns:
            Ответ LLM или None, если релевантных данных нет.
        """
        async def _emit_progress(stage: str, payload: Optional[Dict[str, Any]] = None) -> None:
            """Безопасно отправить событие прогресса во внешний callback."""
            if on_progress is None:
                return
            try:
                await on_progress(stage, payload)
            except Exception as progress_exc:
                logger.warning(
                    "RAG progress callback failed: stage=%s user_id=%s error_type=%s error_repr=%r",
                    stage,
                    user_id,
                    type(progress_exc).__name__,
                    progress_exc,
                )

        normalized_question = (question or "").strip()
        if len(normalized_question) < 3:
            return None

        corpus_version = self._get_corpus_version()
        cache_key = f"{corpus_version}:{normalized_question.lower()}"
        cached = self._answer_cache.get(cache_key)
        now = time.time()
        if cached and cached.expires_at > now:
            return cached.answer

        await _emit_progress(AI_PROGRESS_STAGE_RAG_PREFILTER_STARTED)

        chunks, summary_blocks = self._retrieve_context_for_question(
            normalized_question,
            limit=ai_settings.AI_RAG_TOP_K,
        )
        if not chunks:
            return None

        from src.sbs_helper_telegram_bot.ai_router.llm_provider import get_provider

        context_blocks: List[str] = []
        total_chars = 0
        max_chars = ai_settings.AI_RAG_MAX_CONTEXT_CHARS

        for index, chunk in enumerate(chunks, start=1):
            _, source, chunk_text, _document_id, _chunk_index = self._unpack_chunk_row(chunk)
            block = f"[Блок {index} | {source}]\n{chunk_text}"
            if total_chars + len(block) > max_chars:
                break
            context_blocks.append(block)
            total_chars += len(block)

        provider = get_provider()
        await _emit_progress(
            AI_PROGRESS_STAGE_RAG_AUGMENTED_REQUEST_STARTED,
            {
                "chunks_count": len(context_blocks),
                "summary_blocks_count": len(summary_blocks),
            },
        )
        answer = await provider.chat(
            messages=[{"role": "user", "content": normalized_question}],
            system_prompt=build_rag_prompt(context_blocks, summary_blocks=summary_blocks),
            user_id=user_id,
            purpose="rag_answer",
        )

        self._answer_cache[cache_key] = CachedAnswer(
            answer=answer,
            expires_at=now + self._cache_ttl_seconds,
        )
        self._clear_expired_cache()

        self._log_query(
            user_id=user_id,
            query=normalized_question,
            cache_hit=False,
            chunks_count=len(context_blocks),
        )

        return answer

    def _retrieve_context_for_question(self, question: str, limit: int) -> Tuple[List[Tuple[float, str, str, int]], List[str]]:
        """Собрать релевантные чанки и summary-блоки для RAG-ответа."""
        retrieval_started_at = time.perf_counter()
        tokens = self._tokenize(question)
        if not tokens:
            return [], []

        prefilter_started_at = time.perf_counter()
        prefilter_docs, summary_vector_scores, summary_vector_source = self._prefilter_documents_by_summary(
            question=question,
            question_tokens=tokens,
            limit=ai_settings.AI_RAG_PREFILTER_TOP_DOCS,
        )
        summary_vector_hits = int(self._summary_vector_prefilter_hits)
        prefilter_ms = (time.perf_counter() - prefilter_started_at) * 1000
        prefilter_doc_ids = [doc_id for doc_id, _, _, _ in prefilter_docs]
        base_prefilter_doc_ids = list(prefilter_doc_ids)
        summary_scores = {doc_id: score for doc_id, _, _, score in prefilter_docs}
        normalized_summary_scores = self._build_relative_summary_scores(summary_scores)
        fallback_doc_ids: List[int] = []
        fallback_docs_limit = max(0, int(ai_settings.AI_RAG_SUMMARY_PREFILTER_FALLBACK_DOCS))

        if prefilter_doc_ids and fallback_docs_limit > 0:
            fallback_doc_ids = self._get_fallback_active_document_ids(
                exclude_document_ids=list(prefilter_doc_ids),
                limit=fallback_docs_limit,
            )
            prefilter_doc_ids.extend(fallback_doc_ids)

        prefilter_scope_doc_ids = list(dict.fromkeys(prefilter_doc_ids))

        lexical_started_at = time.perf_counter()
        lexical_chunks = self._search_relevant_chunks(
            question,
            limit=limit,
            prefiltered_doc_ids=prefilter_scope_doc_ids or None,
            summary_scores=summary_scores,
            normalized_summary_scores=normalized_summary_scores,
        )
        lexical_ms = (time.perf_counter() - lexical_started_at) * 1000

        vector_started_at = time.perf_counter()
        vector_chunks = self._search_relevant_chunks_vector(
            question=question,
            prefiltered_doc_ids=prefilter_scope_doc_ids or None,
        )
        vector_ms = (time.perf_counter() - vector_started_at) * 1000

        merge_started_at = time.perf_counter()
        chunks = self._merge_retrieval_candidates(
            lexical_chunks=lexical_chunks,
            vector_chunks=vector_chunks,
            limit=limit,
            summary_scores=summary_scores,
            normalized_summary_scores=normalized_summary_scores,
        )
        merge_ms = (time.perf_counter() - merge_started_at) * 1000
        mode = self._determine_retrieval_mode(
            lexical_chunks=lexical_chunks,
            vector_chunks=vector_chunks,
            selected_chunks=chunks,
        )
        selected_component_scores = self._build_selected_component_scores(
            lexical_chunks=lexical_chunks,
            vector_chunks=vector_chunks,
        )
        lexical_weight = max(0.0, float(ai_settings.AI_RAG_VECTOR_LEXICAL_WEIGHT))
        vector_weight = max(0.0, float(ai_settings.AI_RAG_VECTOR_SEMANTIC_WEIGHT))
        lexical_scorer = ai_settings.get_rag_lexical_scorer()
        selected_unique_docs = len({int(self._unpack_chunk_row(chunk)[3]) for chunk in chunks})
        selected_top_docs = self._build_selected_top_docs_snapshot(chunks)
        top_source = self._format_log_source(str(self._unpack_chunk_row(chunks[0])[1]) if chunks else "none")
        summary_blocks_started_at = time.perf_counter()
        summary_blocks = self._build_summary_blocks(prefilter_docs)
        summary_blocks_ms = (time.perf_counter() - summary_blocks_started_at) * 1000
        retrieval_total_ms = (time.perf_counter() - retrieval_started_at) * 1000
        logger.info(
            "RAG retrieval: mode=%s lexical_scorer=%s tokens=%s prefilter_docs=%s prefilter_scope_docs=%s fallback_docs=%s lexical_hits=%s vector_hits=%s summary_vector_hits=%s summary_vector_source=%s selected=%s selected_unique_docs=%s selected_top_docs=%s top_source=%s timings_ms(total=%.2f prefilter=%.2f lexical=%.2f vector=%.2f merge=%.2f summary_blocks=%.2f)",
            mode,
            lexical_scorer,
            len(tokens),
            len(prefilter_docs),
            len(prefilter_scope_doc_ids),
            len(fallback_doc_ids),
            len(lexical_chunks),
            len(vector_chunks),
            summary_vector_hits,
            summary_vector_source,
            len(chunks),
            selected_unique_docs,
            selected_top_docs,
            top_source,
            retrieval_total_ms,
            prefilter_ms,
            lexical_ms,
            vector_ms,
            merge_ms,
            summary_blocks_ms,
        )
        logger.info(
            "RAG priority evidence:\n  prefilter_top:\n%s\n  selected_top:\n%s",
            self._build_prefilter_priority_snapshot(prefilter_docs, summary_vector_scores),
            self._build_selected_priority_snapshot(
                chunks,
                summary_scores,
                prefilter_scope_doc_ids=prefilter_scope_doc_ids,
                base_prefilter_doc_ids=base_prefilter_doc_ids,
                component_scores=selected_component_scores,
                lexical_weight=lexical_weight,
                vector_weight=vector_weight,
                normalized_summary_scores=normalized_summary_scores,
            ),
        )
        return chunks, summary_blocks

    @staticmethod
    def _format_log_source(source: str, max_length: int = 96) -> str:
        """Подготовить source к компактному человекочитаемому виду для логов."""
        compact = _SPACES_RE.sub(" ", str(source or "").strip())
        if not compact:
            return "none"

        if len(compact) <= max_length:
            return compact

        head_len = max(12, int(max_length * 0.65))
        tail_len = max(8, max_length - head_len - 1)
        return f"{compact[:head_len]}…{compact[-tail_len:]}"

    @classmethod
    def _build_prefilter_priority_snapshot(
        cls,
        prefilter_docs: List[Tuple[int, str, str, float]],
        vector_scores: Optional[Dict[int, float]] = None,
        vector_weight: Optional[float] = None,
    ) -> str:
        """Сформировать человекочитаемый лог prefilter-документов с разложением итогового score."""
        if not prefilter_docs:
            return "    (none)"

        vector_scores = vector_scores or {}
        effective_weight = max(
            0.0,
            float(
                ai_settings.AI_RAG_SUMMARY_VECTOR_WEIGHT if vector_weight is None else vector_weight
            ),
        )
        top_docs = prefilter_docs[:5]
        parts = []
        for rank, (doc_id, filename, _, score) in enumerate(top_docs, start=1):
            vec = float(vector_scores.get(doc_id, 0.0))
            weighted_vec = vec * effective_weight
            lexical_part = score - weighted_vec
            parts.append(
                "    "
                f"{rank}. doc={doc_id} summary={score:.3f} lexical={lexical_part:.3f} "
                f"vec={vec:.3f} vec_w={weighted_vec:.3f} source={cls._format_log_source(filename)}"
            )
        return "\n".join(parts)

    @classmethod
    def _build_selected_priority_snapshot(
        cls,
        chunks: List[Tuple[float, str, str, int]],
        summary_scores: Dict[int, float],
        prefilter_scope_doc_ids: Optional[List[int]] = None,
        base_prefilter_doc_ids: Optional[List[int]] = None,
        component_scores: Optional[Dict[Tuple[int, str], Tuple[float, float]]] = None,
        lexical_weight: Optional[float] = None,
        vector_weight: Optional[float] = None,
        normalized_summary_scores: Optional[Dict[int, float]] = None,
    ) -> str:
        """Сформировать человекочитаемый лог финально выбранных чанков и вкладов summary-score."""
        if not chunks:
            return "    (none)"

        prefilter_scope_set = {int(doc_id) for doc_id in (prefilter_scope_doc_ids or [])}
        base_prefilter_set = {int(doc_id) for doc_id in (base_prefilter_doc_ids or [])}
        component_scores = component_scores or {}
        effective_lexical_weight = max(
            0.0,
            float(ai_settings.AI_RAG_VECTOR_LEXICAL_WEIGHT if lexical_weight is None else lexical_weight),
        )
        effective_vector_weight = max(
            0.0,
            float(ai_settings.AI_RAG_VECTOR_SEMANTIC_WEIGHT if vector_weight is None else vector_weight),
        )
        effective_normalized_scores = (
            normalized_summary_scores
            if normalized_summary_scores is not None
            else cls._build_relative_summary_scores(summary_scores)
        )
        top_chunks = chunks[:5]
        parts = []
        for rank, chunk in enumerate(top_chunks, start=1):
            fused_score, source, chunk_text, document_id, chunk_index = cls._unpack_chunk_row(chunk)
            summary_score = float(summary_scores.get(document_id, 0.0))
            normalized_summary_score = float(effective_normalized_scores.get(document_id, 0.0))
            chunk_key = (int(document_id), str(chunk_text or "").strip())
            lexical_score, vector_score = component_scores.get(chunk_key, (0.0, 0.0))
            lexical_total = float(lexical_score)
            lexical_bonus_full = cls._summary_score_bonus_from_normalized(normalized_summary_score)
            lexical_bonus = max(0.0, min(lexical_bonus_full, lexical_total))
            lexical_raw = max(0.0, lexical_total - lexical_bonus)
            hybrid_base = (float(lexical_score) * effective_lexical_weight) + (float(vector_score) * effective_vector_weight)
            summary_bonus = cls._summary_postrank_bonus_from_normalized(normalized_summary_score)
            if document_id in base_prefilter_set:
                origin = "prefilter"
            elif document_id in prefilter_scope_set:
                origin = "fallback"
            else:
                origin = "global"
            parts.append(
                "    "
                f"{rank}. doc={document_id} chunk={chunk_index} fused={float(fused_score):.3f} "
                f"summary={summary_score:.3f} origin={origin} "
                f"lex_raw={lexical_raw:.3f} lex_bonus={lexical_bonus:.3f} lex_total={lexical_total:.3f} "
                f"hybrid=({lexical_total:.3f}*{effective_lexical_weight:.3f})+({float(vector_score):.3f}*{effective_vector_weight:.3f})={hybrid_base:.3f} "
                f"summary_bonus={summary_bonus:.3f} source={cls._format_log_source(source)}"
            )
        return "\n".join(parts)

    @staticmethod
    def _build_selected_component_scores(
        lexical_chunks: List[Tuple[float, str, str, int]],
        vector_chunks: List[Tuple[float, str, str, int]],
    ) -> Dict[Tuple[int, str], Tuple[float, float]]:
        """Собрать lexical/vector score-компоненты для выбранных чанков по dedup-ключу merge."""
        components: Dict[Tuple[int, str], Tuple[float, float]] = {}

        for chunk in lexical_chunks:
            lexical_score, _source, chunk_text, document_id, _chunk_index = RagKnowledgeService._unpack_chunk_row(chunk)
            key = (int(document_id), str(chunk_text or "").strip())
            current_lexical, current_vector = components.get(key, (0.0, 0.0))
            components[key] = (max(current_lexical, float(lexical_score)), current_vector)

        for chunk in vector_chunks:
            vector_score, _source, chunk_text, document_id, _chunk_index = RagKnowledgeService._unpack_chunk_row(chunk)
            key = (int(document_id), str(chunk_text or "").strip())
            current_lexical, current_vector = components.get(key, (0.0, 0.0))
            components[key] = (current_lexical, max(current_vector, float(vector_score)))

        return components

    @staticmethod
    def _build_selected_top_docs_snapshot(
        chunks: List[Tuple[float, str, str, int]],
        max_docs: int = 5,
    ) -> str:
        """Сформировать compact snapshot top уникальных document_id по порядку ранжирования."""
        if not chunks:
            return "none"

        safe_limit = max(1, int(max_docs))
        ordered_unique_doc_ids: List[int] = []
        seen_doc_ids: set[int] = set()
        for chunk in chunks:
            _score, _source, _chunk_text, document_id, _chunk_index = RagKnowledgeService._unpack_chunk_row(chunk)
            doc_id = int(document_id)
            if doc_id in seen_doc_ids:
                continue
            seen_doc_ids.add(doc_id)
            ordered_unique_doc_ids.append(doc_id)
            if len(ordered_unique_doc_ids) >= safe_limit:
                break

        if not ordered_unique_doc_ids:
            return "none"
        return ",".join(str(doc_id) for doc_id in ordered_unique_doc_ids)

    def _search_relevant_chunks_vector(
        self,
        question: str,
        prefiltered_doc_ids: Optional[List[int]],
    ) -> List[Tuple[float, str, str, int, int]]:
        """Найти релевантные чанки через локальный векторный индекс."""
        if not self._is_vector_search_enabled():
            return []

        embedding_provider = self._get_embedding_provider()
        vector_index = self._get_vector_index()
        if embedding_provider is None or vector_index is None:
            return []

        query_vectors = embedding_provider.encode_texts([question])
        if not query_vectors:
            return []

        candidates = vector_index.search(
            query_vector=query_vectors[0],
            limit=max(1, int(ai_settings.AI_RAG_VECTOR_TOP_K)),
            allowed_document_ids=prefiltered_doc_ids,
        )
        return [(item.score, item.source, item.chunk_text, item.document_id, item.chunk_index) for item in candidates]

    def _merge_retrieval_candidates(
        self,
        lexical_chunks: List[Tuple[float, str, str, int]],
        vector_chunks: List[Tuple[float, str, str, int]],
        limit: int,
        summary_scores: Optional[Dict[int, float]] = None,
        normalized_summary_scores: Optional[Dict[int, float]] = None,
    ) -> List[Tuple[float, str, str, int, int]]:
        """Объединить lexical и vector кандидаты в единый ранжированный список."""
        safe_limit = max(1, limit)
        summary_scores = summary_scores or {}
        normalized_summary_scores = normalized_summary_scores or {}

        if not vector_chunks:
            return lexical_chunks[:safe_limit]

        if not ai_settings.AI_RAG_HYBRID_ENABLED:
            return vector_chunks[:safe_limit]

        lexical_weight = max(0.0, float(ai_settings.AI_RAG_VECTOR_LEXICAL_WEIGHT))
        vector_weight = max(0.0, float(ai_settings.AI_RAG_VECTOR_SEMANTIC_WEIGHT))

        merged: Dict[Tuple[int, int, str], Dict[str, object]] = {}

        for chunk in lexical_chunks:
            score, source, chunk_text, document_id, chunk_index = self._unpack_chunk_row(chunk)
            key = self._chunk_merge_key(document_id=document_id, chunk_text=chunk_text, chunk_index=chunk_index)
            merged[key] = {
                "lexical_score": float(score),
                "vector_score": 0.0,
                "source": source,
                "chunk_text": chunk_text,
                "document_id": document_id,
                "chunk_index": int(chunk_index),
            }

        for chunk in vector_chunks:
            score, source, chunk_text, document_id, chunk_index = self._unpack_chunk_row(chunk)
            key = self._chunk_merge_key(document_id=document_id, chunk_text=chunk_text, chunk_index=chunk_index)
            row = merged.get(key)
            if row is None:
                row = {
                    "lexical_score": 0.0,
                    "vector_score": 0.0,
                    "source": source,
                    "chunk_text": chunk_text,
                    "document_id": document_id,
                    "chunk_index": int(chunk_index),
                }
                merged[key] = row
            row["vector_score"] = max(float(row.get("vector_score") or 0.0), float(score))

        ranked: List[Tuple[float, str, str, int, int]] = []
        for row in merged.values():
            lexical_score = float(row.get("lexical_score") or 0.0)
            vector_score = float(row.get("vector_score") or 0.0)
            document_id = int(row.get("document_id") or 0)
            chunk_index = int(row.get("chunk_index") or 0)
            summary_score = float(summary_scores.get(document_id, 0.0))
            normalized_summary_score = normalized_summary_scores.get(document_id)
            fused_score = (lexical_score * lexical_weight) + (vector_score * vector_weight)
            if normalized_summary_score is None:
                fused_score += self._summary_postrank_bonus(summary_score)
            else:
                fused_score += self._summary_postrank_bonus_from_normalized(float(normalized_summary_score))
            if fused_score <= 0:
                continue
            ranked.append(
                (
                    fused_score,
                    str(row.get("source") or "document"),
                    str(row.get("chunk_text") or ""),
                    int(row.get("document_id") or 0),
                    chunk_index,
                )
            )

        ranked.sort(key=lambda item: item[0], reverse=True)
        return ranked[:safe_limit]

    def _extract_text(self, filename: str, payload: bytes) -> str:
        """Извлечь текст из поддерживаемого формата документа."""
        lower_name = filename.lower()

        if lower_name.endswith(".txt") or lower_name.endswith(".md"):
            return self._decode_text_payload(payload)

        if lower_name.endswith(".html") or lower_name.endswith(".htm"):
            return self._extract_html_text(payload)

        if lower_name.endswith(".pdf"):
            return self._extract_pdf_text(payload)

        if lower_name.endswith(".docx"):
            return self._extract_docx_text(payload)

        raise ValueError("Неподдерживаемый формат файла")

    @staticmethod
    def _is_html_file(filename: str) -> bool:
        """Проверить, относится ли имя файла к HTML-формату."""
        lower_name = (filename or "").lower()
        return lower_name.endswith(".html") or lower_name.endswith(".htm")

    @staticmethod
    def _decode_text_payload(payload: bytes) -> str:
        """Декодировать текстовый файл с fallback по кодировкам."""
        for encoding in ("utf-8", "cp1251", "latin-1"):
            try:
                return payload.decode(encoding)
            except UnicodeDecodeError:
                continue
        return payload.decode("utf-8", errors="ignore")

    @staticmethod
    def _extract_pdf_text(payload: bytes) -> str:
        """Извлечь текст из PDF документа."""
        try:
            from pypdf import PdfReader
        except Exception as exc:
            raise ValueError("Для PDF требуется пакет pypdf") from exc

        reader = PdfReader(io.BytesIO(payload))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    @staticmethod
    def _extract_docx_text(payload: bytes) -> str:
        """Извлечь текст из DOCX документа."""
        try:
            from docx import Document
        except Exception as exc:
            raise ValueError("Для DOCX требуется пакет python-docx") from exc

        doc = Document(io.BytesIO(payload))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)

    @staticmethod
    def _extract_html_text(payload: bytes) -> str:
        """Извлечь текст из HTML документа."""
        raw_html = RagKnowledgeService._decode_text_payload(payload)

        no_script = re.sub(r"<script[\s\S]*?</script>", " ", raw_html, flags=re.IGNORECASE)
        no_style = re.sub(r"<style[\s\S]*?</style>", " ", no_script, flags=re.IGNORECASE)
        no_tags = re.sub(r"<[^>]+>", " ", no_style)
        compact = re.sub(r"[ \t\r\f\v]+", " ", no_tags)
        compact = re.sub(r"\n\s*\n+", "\n", compact)
        return compact.strip()

    def _split_html_payload(self, payload: bytes, filename: str = "document.html") -> List[str]:
        """Разбить HTML-документ на чанки с приоритетом по заголовкам."""
        if not ai_settings.is_rag_html_splitter_enabled():
            logger.info("HTML splitter отключен через bot_settings, используется fallback")
            chunks = self._split_text(self._extract_html_text(payload))
            self._log_chunking_strategy(
                file_name=filename,
                file_format="html",
                strategy="plain_text_fallback(html_splitter_disabled)",
                chunks_count=len(chunks),
            )
            return chunks

        raw_html = self._decode_text_payload(payload)
        semantic_chunks = self._split_html_with_semantic_preserving_splitter(raw_html)
        if semantic_chunks:
            self._log_chunking_strategy(
                file_name=filename,
                file_format="html",
                strategy="html_semantic_preserving_splitter",
                chunks_count=len(semantic_chunks),
            )
            return semantic_chunks

        logger.info("HTMLSemanticPreservingSplitter/HTMLHeaderTextSplitter недоступны или не дали чанков, включен fallback")
        chunks = self._split_text(self._extract_html_text(payload))
        self._log_chunking_strategy(
            file_name=filename,
            file_format="html",
            strategy="plain_text_fallback(html_splitter_empty_or_unavailable)",
            chunks_count=len(chunks),
        )
        return chunks

    def _split_html_with_semantic_preserving_splitter(self, raw_html: str) -> List[str]:
        """Попытаться разбить HTML через semantic-preserving splitter с переносом заголовков в текст."""
        normalized_html = (raw_html or "").strip()
        if not normalized_html:
            return []

        try:
            splitter_cls = self._get_html_splitter_class()
            splitter = self._build_html_splitter(splitter_cls)
            documents = splitter.split_text(normalized_html)
        except Exception as exc:
            logger.warning("Не удалось применить HTML splitter: %s", exc)
            return []

        chunks: List[str] = []
        header_keys = ("h1", "h2", "h3", "h4", "h5", "h6")

        for doc in documents or []:
            page_content = str(getattr(doc, "page_content", "") or "").strip()
            metadata = getattr(doc, "metadata", {}) or {}
            metadata_values = [
                str(metadata.get(key, "")).strip() for key in header_keys if str(metadata.get(key, "")).strip()
            ]

            if not page_content and not metadata_values:
                continue

            combined_block = "\n".join(metadata_values + [page_content]).strip()
            if not combined_block:
                continue

            chunks.extend(self._split_text(combined_block))

        return [chunk for chunk in chunks if chunk.strip()]

    @staticmethod
    def _build_html_splitter(splitter_cls):
        """Построить HTML splitter с совместимыми параметрами конструктора."""
        headers_to_split_on = [
            ("h1", "h1"),
            ("h2", "h2"),
            ("h3", "h3"),
            ("h4", "h4"),
            ("h5", "h5"),
            ("h6", "h6"),
        ]
        try:
            return splitter_cls(
                headers_to_split_on=headers_to_split_on,
                max_chunk_size=ai_settings.AI_RAG_CHUNK_SIZE,
                chunk_overlap=ai_settings.AI_RAG_CHUNK_OVERLAP,
            )
        except TypeError:
            return splitter_cls(headers_to_split_on=headers_to_split_on)

    @staticmethod
    def _get_html_splitter_class():
        """Получить приоритетный HTML splitter: semantic-preserving, затем header splitter."""
        if not RagKnowledgeService._is_langchain_splitter_supported():
            raise RuntimeError("LangChain splitter отключен для текущей версии Python")

        import warnings

        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                from langchain_text_splitters import HTMLSemanticPreservingSplitter

            return HTMLSemanticPreservingSplitter
        except Exception:
            pass

        return RagKnowledgeService._get_html_header_splitter_class()

    @staticmethod
    def _get_html_header_splitter_class():
        """Получить fallback-класс HTMLHeaderTextSplitter из доступного пространства имён."""
        import warnings

        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                from langchain_text_splitters import HTMLHeaderTextSplitter

            return HTMLHeaderTextSplitter
        except Exception:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                from langchain.text_splitter import HTMLHeaderTextSplitter

            return HTMLHeaderTextSplitter

    @staticmethod
    def _is_langchain_splitter_supported() -> bool:
        """Проверить, можно ли безопасно использовать LangChain splitters в текущем Python."""
        try:
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                from langchain_text_splitters import HTMLSemanticPreservingSplitter  # noqa: F401
            return True
        except Exception:
            pass
        try:
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                from langchain_text_splitters import HTMLHeaderTextSplitter  # noqa: F401
            return True
        except Exception:
            pass
        try:
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                from langchain.text_splitter import HTMLHeaderTextSplitter  # noqa: F401
            return True
        except Exception:
            logger.info(
                "LangChain splitters недоступны на Python %s.%s: используется fallback chunking",
                sys.version_info.major,
                sys.version_info.minor,
            )
            return False

    @staticmethod
    def _split_text(text: str) -> List[str]:
        """Разбить текст на чанки; при наличии langchain использует его splitter."""
        cleaned = (text or "").strip()
        if not cleaned:
            return []

        if RagKnowledgeService._is_langchain_splitter_supported():
            try:
                from langchain.text_splitter import RecursiveCharacterTextSplitter

                splitter = RecursiveCharacterTextSplitter(
                    chunk_size=ai_settings.AI_RAG_CHUNK_SIZE,
                    chunk_overlap=ai_settings.AI_RAG_CHUNK_OVERLAP,
                    separators=["\n\n", "\n", ". ", " ", ""],
                )
                chunks = splitter.split_text(cleaned)
                return [chunk.strip() for chunk in chunks if chunk and chunk.strip()]
            except Exception:
                pass

        chunk_size = ai_settings.AI_RAG_CHUNK_SIZE
        overlap = ai_settings.AI_RAG_CHUNK_OVERLAP
        chunks: List[str] = []
        start = 0
        text_len = len(cleaned)

        while start < text_len:
            end = min(start + chunk_size, text_len)
            chunk = cleaned[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= text_len:
                break
            start = max(end - overlap, start + 1)

        return chunks

    def _search_relevant_chunks(
        self,
        question: str,
        limit: int,
        prefiltered_doc_ids: Optional[List[int]] = None,
        summary_scores: Optional[Dict[int, float]] = None,
        normalized_summary_scores: Optional[Dict[int, float]] = None,
    ) -> List[Tuple[float, str, str, int, int]]:
        """Найти релевантные чанки в БД по гибридному lexical scoring."""
        tokens = self._tokenize(question)
        if not tokens:
            return []

        lexical_scorer = ai_settings.get_rag_lexical_scorer()

        safe_limit = max(1, limit)
        summary_scores = summary_scores or {}
        normalized_summary_scores = normalized_summary_scores or {}

        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                if prefiltered_doc_ids:
                    placeholders = ",".join(["%s"] * len(prefiltered_doc_ids))
                    cursor.execute(
                        f"""
                        SELECT c.chunk_text, c.chunk_index, d.filename, d.id AS document_id
                        FROM rag_chunks c
                        JOIN rag_documents d ON d.id = c.document_id
                        WHERE d.status = 'active' AND d.id IN ({placeholders})
                        ORDER BY c.id DESC
                        LIMIT %s
                        """,
                        tuple(prefiltered_doc_ids) + (_RAG_CHUNK_SCAN_LIMIT,),
                    )
                else:
                    cursor.execute(
                        """
                        SELECT c.chunk_text, c.chunk_index, d.filename, d.id AS document_id
                        FROM rag_chunks c
                        JOIN rag_documents d ON d.id = c.document_id
                        WHERE d.status = 'active'
                        ORDER BY c.id DESC
                        LIMIT %s
                        """,
                        (_RAG_CHUNK_SCAN_LIMIT,),
                    )
                rows = cursor.fetchall() or []

        scored: List[Tuple[float, str, str, int, int]] = []

        bm25_scores: List[float] = []
        if lexical_scorer == "bm25":
            chunk_tokens_corpus = [self._tokenize(str(row.get("chunk_text") or "")) for row in rows]
            bm25_scores = self._score_corpus_bm25(chunk_tokens_corpus, tokens)

        for index, row in enumerate(rows):
            chunk_text = row.get("chunk_text") or ""
            chunk_index = int(row.get("chunk_index") or 0)
            source = row.get("filename") or "document"
            document_id = int(row.get("document_id") or 0)
            if lexical_scorer == "bm25":
                chunk_score = float(bm25_scores[index]) if index < len(bm25_scores) else 0.0
            else:
                chunk_score = self._score_chunk(chunk_text, tokens)
            summary_score = summary_scores.get(document_id, 0.0)
            normalized_summary_score = normalized_summary_scores.get(document_id)
            if normalized_summary_score is None:
                summary_bonus = self._summary_score_bonus(summary_score)
            else:
                summary_bonus = self._summary_score_bonus_from_normalized(float(normalized_summary_score))
            score = chunk_score + summary_bonus
            if score > 0:
                scored.append((score, source, chunk_text, document_id, chunk_index))

        scored.sort(key=lambda item: item[0], reverse=True)
        return scored[:safe_limit]

    @staticmethod
    def _unpack_chunk_row(chunk: Tuple[object, ...]) -> Tuple[float, str, str, int, int]:
        """Нормализовать запись чанка к формату (score, source, text, document_id, chunk_index)."""
        if len(chunk) >= 5:
            score, source, chunk_text, document_id, chunk_index = chunk[:5]
            return (
                float(score or 0.0),
                str(source or "document"),
                str(chunk_text or ""),
                int(document_id or 0),
                int(chunk_index or 0),
            )

        if len(chunk) == 4:
            score, source, chunk_text, document_id = chunk
            return (
                float(score or 0.0),
                str(source or "document"),
                str(chunk_text or ""),
                int(document_id or 0),
                0,
            )

        raise ValueError("Некорректный формат чанка для retrieval")

    @staticmethod
    def _chunk_merge_key(document_id: int, chunk_text: str, chunk_index: int) -> Tuple[int, int, str]:
        """Собрать dedup-ключ merge с приоритетом chunk_index при его наличии."""
        safe_doc_id = int(document_id or 0)
        safe_chunk_index = int(chunk_index or 0)
        normalized_text = str(chunk_text or "").strip()
        if safe_chunk_index > 0:
            return safe_doc_id, safe_chunk_index, ""
        return safe_doc_id, 0, normalized_text

    def _prefilter_documents_by_summary(
        self,
        question: str,
        question_tokens: List[str],
        limit: int,
    ) -> Tuple[List[Tuple[int, str, str, float]], Dict[int, float], str]:
        """Отобрать релевантные документы по таблице summary перед поиском чанков.

        Returns:
            (список ранжированных документов, словарь vector-similarity по document_id, источник vector-score)
        """
        if not question_tokens:
            self._summary_vector_prefilter_source = "disabled"
            self._summary_vector_prefilter_hits = 0
            return [], {}, self._summary_vector_prefilter_source

        safe_limit = max(1, min(limit, 100))
        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                cursor.execute(
                    """
                    SELECT s.document_id, s.summary_text, d.filename
                    FROM rag_document_summaries s
                    JOIN rag_documents d ON d.id = s.document_id
                    WHERE d.status = 'active'
                    ORDER BY s.updated_at DESC
                    LIMIT %s
                    """,
                    (_RAG_SUMMARY_SCAN_LIMIT,),
                )
                rows = cursor.fetchall() or []

        summaries_for_vector: List[Tuple[int, str]] = []
        valid_rows: List[Tuple[int, str, str]] = []
        for row in rows:
            document_id = int(row.get("document_id") or 0)
            filename = str(row.get("filename") or "document")
            summary_text = str(row.get("summary_text") or "").strip()
            if not document_id or not summary_text:
                continue
            valid_rows.append((document_id, filename, summary_text))
            summaries_for_vector.append((document_id, summary_text))

        vector_scores = self._search_summary_vector_scores_from_collection(
            question=question,
            document_ids=[doc_id for doc_id, _, _ in valid_rows],
            limit=max(safe_limit, len(valid_rows)),
        )
        vector_source = "collection"
        if not vector_scores:
            vector_scores = self._compute_summary_vector_scores(
                question=question,
                summaries=summaries_for_vector,
            )
            vector_source = "fallback"
        self._summary_vector_prefilter_source = vector_source
        self._summary_vector_prefilter_hits = len(vector_scores)

        vector_weight = max(0.0, float(ai_settings.AI_RAG_SUMMARY_VECTOR_WEIGHT))
        lexical_scorer = ai_settings.get_rag_lexical_scorer()
        summary_bm25_scores: Dict[int, float] = {}
        if lexical_scorer == "bm25":
            summary_corpus_tokens = [self._tokenize(summary_text) for _, _, summary_text in valid_rows]
            corpus_scores = self._score_corpus_bm25(summary_corpus_tokens, question_tokens)
            for index, (document_id, _, _) in enumerate(valid_rows):
                summary_bm25_scores[document_id] = float(corpus_scores[index]) if index < len(corpus_scores) else 0.0

        token_weight = max(0.0, float(ai_settings.AI_RAG_SUMMARY_MATCH_TOKEN_WEIGHT))

        scored_docs: List[Tuple[int, str, str, float]] = []
        for document_id, filename, summary_text in valid_rows:
            if lexical_scorer == "bm25":
                lexical_score = summary_bm25_scores.get(document_id, 0.0) * token_weight
                lexical_score += self._score_summary_phrase_match(
                    summary_text=summary_text,
                    question_tokens=question_tokens,
                    question=question,
                )
            else:
                lexical_score = self._score_summary_text(
                    summary_text=summary_text,
                    question_tokens=question_tokens,
                    question=question,
                )
            vec_score = float(vector_scores.get(document_id, 0.0))
            score = lexical_score + (vec_score * vector_weight)
            if score <= 0:
                continue
            scored_docs.append((document_id, filename, summary_text, score))

        scored_docs.sort(key=lambda item: item[3], reverse=True)
        return scored_docs[:safe_limit], vector_scores, vector_source

    def _search_summary_vector_scores_from_collection(
        self,
        question: str,
        document_ids: List[int],
        limit: int,
    ) -> Dict[int, float]:
        """Получить summary vector-score из отдельной Qdrant-коллекции по списку document_id."""
        if not ai_settings.is_rag_summary_vector_enabled() or not self._is_vector_search_enabled():
            return {}

        embedding_provider = self._get_embedding_provider()
        vector_index = self._get_vector_index()
        if embedding_provider is None or vector_index is None:
            return {}

        query_vectors = embedding_provider.encode_texts([question])
        if not query_vectors:
            return {}

        candidates = vector_index.search_summaries(
            query_vector=query_vectors[0],
            limit=max(1, int(limit), int(ai_settings.get_rag_summary_vector_top_k())),
            allowed_document_ids=document_ids,
        )
        if not candidates:
            return {}

        scores: Dict[int, float] = {}
        for candidate in candidates:
            safe_doc_id = int(candidate.document_id or 0)
            if safe_doc_id <= 0:
                continue
            safe_score = max(0.0, float(candidate.score or 0.0))
            scores[safe_doc_id] = max(safe_score, scores.get(safe_doc_id, 0.0))

        return scores

    @staticmethod
    def _build_summary_blocks(prefilter_docs: List[Tuple[int, str, str, float]]) -> List[str]:
        """Собрать summary-блоки для системного RAG-промпта."""
        max_docs = max(0, int(ai_settings.AI_RAG_PROMPT_SUMMARY_DOCS))
        if max_docs <= 0:
            return []

        summary_blocks: List[str] = []
        for _, filename, summary_text, _ in prefilter_docs[:max_docs]:
            safe_summary = summary_text.strip()
            if not safe_summary:
                continue
            summary_blocks.append(f"[Summary | {filename}]\n{safe_summary}")
        return summary_blocks

    @staticmethod
    def _summary_score_bonus(summary_score: float) -> float:
        """Рассчитать бонус чанку на основе релевантности summary документа."""
        normalized_summary_score = RagKnowledgeService._normalize_summary_score(summary_score)
        return RagKnowledgeService._summary_score_bonus_from_normalized(normalized_summary_score)

    @staticmethod
    def _summary_score_bonus_from_normalized(normalized_summary_score: float) -> float:
        """Рассчитать бонус чанку из нормализованного summary-score (0..1)."""
        safe_score = max(0.0, min(1.0, float(normalized_summary_score or 0.0)))
        if safe_score <= 0:
            return 0.0

        bonus_weight = max(0.0, float(ai_settings.AI_RAG_SUMMARY_BONUS_WEIGHT))
        return safe_score * bonus_weight

    @staticmethod
    def _summary_postrank_bonus(summary_score: float) -> float:
        """Рассчитать пост-бонус документа при финальном hybrid-ранжировании."""
        normalized_summary_score = RagKnowledgeService._normalize_summary_score(summary_score)
        return RagKnowledgeService._summary_postrank_bonus_from_normalized(normalized_summary_score)

    @staticmethod
    def _summary_postrank_bonus_from_normalized(normalized_summary_score: float) -> float:
        """Рассчитать пост-бонус документа из нормализованного summary-score (0..1)."""
        safe_score = max(0.0, min(1.0, float(normalized_summary_score or 0.0)))
        if safe_score <= 0:
            return 0.0

        postrank_weight = max(0.0, float(ai_settings.AI_RAG_SUMMARY_POSTRANK_WEIGHT))
        return safe_score * postrank_weight

    @staticmethod
    def _normalize_summary_score(summary_score: float) -> float:
        """Нормализовать summary-score документа в диапазон 0..1 по верхней границе cap."""
        if summary_score <= 0:
            return 0.0

        score_cap = max(0.0, float(ai_settings.AI_RAG_SUMMARY_SCORE_CAP))
        if score_cap <= 0:
            return 0.0

        return min(summary_score, score_cap) / score_cap

    @staticmethod
    def _build_relative_summary_scores(summary_scores: Dict[int, float]) -> Dict[int, float]:
        """Нормализовать summary-score документов в текущем пуле retrieval по min-max в диапазон 0..1."""
        if not summary_scores:
            return {}

        cleaned_scores: Dict[int, float] = {}
        for document_id, score in summary_scores.items():
            safe_doc_id = int(document_id or 0)
            safe_score = max(0.0, float(score or 0.0))
            if safe_doc_id <= 0 or safe_score <= 0:
                continue
            cleaned_scores[safe_doc_id] = safe_score

        if not cleaned_scores:
            return {}

        min_score = min(cleaned_scores.values())
        max_score = max(cleaned_scores.values())
        if max_score <= min_score:
            return {doc_id: 1.0 for doc_id in cleaned_scores}

        denominator = max_score - min_score
        normalized_scores: Dict[int, float] = {}
        for doc_id, score in cleaned_scores.items():
            normalized = (score - min_score) / denominator
            normalized_scores[doc_id] = max(0.0, min(1.0, normalized))

        return normalized_scores

    def _score_summary_text(self, summary_text: str, question_tokens: List[str], question: str) -> float:
        """Оценить релевантность summary по фразовому и токенному совпадению.

        Фразовый матч использует word-boundary regex, чтобы 'X5' не
        совпадало внутри 'VX520'.
        """
        low_summary = self._normalize_for_phrase_match(summary_text)
        if not low_summary:
            return 0.0

        token_weight = max(0.0, float(ai_settings.AI_RAG_SUMMARY_MATCH_TOKEN_WEIGHT))
        token_score = self._score_chunk(summary_text, question_tokens)
        phrase_score = self._score_summary_phrase_match(
            summary_text=summary_text,
            question_tokens=question_tokens,
            question=question,
        )

        return (token_score * token_weight) + phrase_score

    def _score_summary_phrase_match(self, summary_text: str, question_tokens: List[str], question: str) -> float:
        """Оценить фразовое совпадение вопроса и summary (с учётом word-boundary)."""
        low_summary = self._normalize_for_phrase_match(summary_text)
        if not low_summary:
            return 0.0

        phrase_weight = max(0.0, float(ai_settings.AI_RAG_SUMMARY_MATCH_PHRASE_WEIGHT))
        phrase_score = 0.0

        normalized_question = self._normalize_for_phrase_match(question)
        normalized_token_phrase = self._normalize_for_phrase_match(" ".join(question_tokens))

        if len(normalized_question) >= 2 and self._word_boundary_match(normalized_question, low_summary):
            phrase_score = 1.0
        elif normalized_token_phrase and self._word_boundary_match(normalized_token_phrase, low_summary):
            phrase_score = 0.8

        return phrase_score * phrase_weight

    @staticmethod
    def _normalize_for_phrase_match(text: str) -> str:
        """Нормализовать текст для фразового матчингa без учёта регистра и лишних пробелов."""
        normalized = _SPACES_RE.sub(" ", (text or "").lower())
        return normalized.strip()

    @staticmethod
    def _word_boundary_match(phrase: str, text: str) -> bool:
        """Проверить наличие фразы в тексте по word-boundary (\\b).

        Используется вместо оператора ``in``, чтобы ``'x5'`` не
        совпадало внутри ``'vx520'``.
        """
        if not phrase or not text:
            return False
        try:
            return bool(re.search(r"\b" + re.escape(phrase) + r"\b", text, re.IGNORECASE))
        except re.error:
            return phrase in text

    def _compute_summary_vector_scores(
        self,
        question: str,
        summaries: List[Tuple[int, str]],
    ) -> Dict[int, float]:
        """Вычислить семантическое сходство вопроса и каждого summary через эмбеддинги.

        Результат — словарь {document_id: cosine_similarity}.  Если
        embedding-провайдер недоступен, возвращается пустой словарь
        (graceful degradation).
        """
        if not summaries:
            return {}

        embedding_provider = self._get_embedding_provider()
        if embedding_provider is None:
            return {}

        corpus_version = self._get_corpus_version()
        if corpus_version != self._summary_embedding_corpus_version:
            self._summary_embedding_cache.clear()
            self._summary_embedding_corpus_version = corpus_version

        question_vectors = embedding_provider.encode_texts([question])
        if not question_vectors:
            return {}
        q_vec = question_vectors[0]

        uncached_ids: List[int] = []
        uncached_texts: List[str] = []
        for doc_id, summary_text in summaries:
            if doc_id not in self._summary_embedding_cache:
                uncached_ids.append(doc_id)
                uncached_texts.append(summary_text)

        if uncached_texts:
            new_vectors = embedding_provider.encode_texts(uncached_texts)
            for idx, doc_id in enumerate(uncached_ids):
                if idx < len(new_vectors):
                    self._summary_embedding_cache[doc_id] = new_vectors[idx]

        scores: Dict[int, float] = {}
        for doc_id, _ in summaries:
            s_vec = self._summary_embedding_cache.get(doc_id)
            if s_vec is None:
                continue
            similarity = self._cosine_dot(q_vec, s_vec)
            scores[doc_id] = max(0.0, similarity)

        return scores

    @staticmethod
    def _cosine_dot(vec_a: List[float], vec_b: List[float]) -> float:
        """Вычислить скалярное произведение (≈cosine similarity для L2-нормированных)."""
        return sum(a * b for a, b in zip(vec_a, vec_b))

    @staticmethod
    def _get_fallback_active_document_ids(
        exclude_document_ids: List[int],
        limit: int,
    ) -> List[int]:
        """Получить fallback-список активных документов для сохранения recall."""
        safe_limit = max(0, min(int(limit), 100))
        if safe_limit <= 0:
            return []

        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                if exclude_document_ids:
                    placeholders = ",".join(["%s"] * len(exclude_document_ids))
                    cursor.execute(
                        f"""
                        SELECT id
                        FROM rag_documents
                        WHERE status = 'active' AND id NOT IN ({placeholders})
                        ORDER BY updated_at DESC, id DESC
                        LIMIT %s
                        """,
                        tuple(exclude_document_ids) + (safe_limit,),
                    )
                else:
                    cursor.execute(
                        """
                        SELECT id
                        FROM rag_documents
                        WHERE status = 'active'
                        ORDER BY updated_at DESC, id DESC
                        LIMIT %s
                        """,
                        (safe_limit,),
                    )
                rows = cursor.fetchall() or []

        fallback_ids: List[int] = []
        for row in rows:
            document_id = int(row.get("id") or 0)
            if document_id > 0:
                fallback_ids.append(document_id)
        return fallback_ids

    def _determine_retrieval_mode(
        self,
        lexical_chunks: List[Tuple[float, str, str, int]],
        vector_chunks: List[Tuple[float, str, str, int]],
        selected_chunks: List[Tuple[float, str, str, int]],
    ) -> str:
        """Определить режим retrieval для диагностического логирования."""
        if not selected_chunks:
            return "lexical_only"

        has_lexical = len(lexical_chunks) > 0
        has_vector = len(vector_chunks) > 0

        if has_lexical and has_vector and ai_settings.AI_RAG_HYBRID_ENABLED:
            return "hybrid"

        if has_vector and (not ai_settings.AI_RAG_HYBRID_ENABLED or not has_lexical):
            return "vector_only"

        if has_lexical and not has_vector and ai_settings.AI_RAG_HYBRID_ENABLED and self._is_vector_search_enabled():
            return "lexical_fallback"

        return "lexical_only"

    def _is_vector_search_enabled(self) -> bool:
        """Проверить, включён ли векторный retrieval через конфигурацию."""
        return bool(ai_settings.AI_RAG_VECTOR_ENABLED)

    def _get_embedding_provider(self) -> Optional[LocalEmbeddingProvider]:
        """Получить lazy-singleton провайдера локальных эмбеддингов."""
        if not self._is_vector_search_enabled():
            return None

        if self._embedding_provider is None:
            self._embedding_provider = LocalEmbeddingProvider()

        if not self._embedding_provider.is_ready():
            return None
        return self._embedding_provider

    def _get_vector_index(self) -> Optional[LocalVectorIndex]:
        """Получить lazy-singleton локального векторного индекса."""
        if not self._is_vector_search_enabled():
            return None

        if self._vector_index is None:
            self._vector_index = LocalVectorIndex()

        if not self._vector_index.is_ready():
            return None

        return self._vector_index

    def _upsert_vectors_for_chunks(self, chunks: List[Dict[str, object]]) -> int:
        """Записать эмбеддинги чанков в локальный векторный индекс."""
        if not chunks or not self._is_vector_search_enabled():
            return 0

        embedding_provider = self._get_embedding_provider()
        vector_index = self._get_vector_index()
        if embedding_provider is None or vector_index is None:
            return 0

        texts = [str(chunk.get("chunk_text") or "") for chunk in chunks]
        embeddings = embedding_provider.encode_texts(texts)
        if not embeddings:
            self._record_chunk_embedding_metadata(
                chunks=chunks,
                embeddings=[],
                status="failed",
                error_message="embedding_unavailable",
            )
            return 0

        upsert_started_at = time.perf_counter()
        upserted = vector_index.upsert_chunks(chunks=chunks, embeddings=embeddings)
        upsert_duration_ms = (time.perf_counter() - upsert_started_at) * 1000
        if upserted > 0:
            self._record_chunk_embedding_metadata(
                chunks=chunks,
                embeddings=embeddings,
                status="ready",
                error_message=None,
            )
        else:
            self._record_chunk_embedding_metadata(
                chunks=chunks,
                embeddings=[],
                status="failed",
                error_message="vector_upsert_failed",
            )
        if upserted > 0:
            logger.info("RAG vector upsert: chunks=%s duration_ms=%.2f", upserted, upsert_duration_ms)
        return upserted

    def _upsert_vectors_for_summaries(self, summaries: List[Dict[str, object]]) -> int:
        """Записать эмбеддинги summary-документов в отдельную vector-коллекцию."""
        if not summaries or not self._is_vector_search_enabled() or not ai_settings.is_rag_summary_vector_enabled():
            return 0

        embedding_provider = self._get_embedding_provider()
        vector_index = self._get_vector_index()
        if embedding_provider is None or vector_index is None:
            return 0

        texts = [str(summary.get("summary_text") or "") for summary in summaries]
        embeddings = embedding_provider.encode_texts(texts)
        if not embeddings:
            self._record_summary_embedding_metadata(
                summaries=summaries,
                embeddings=[],
                status="failed",
                error_message="embedding_unavailable",
            )
            return 0

        upsert_started_at = time.perf_counter()
        upserted = vector_index.upsert_summaries(summaries=summaries, embeddings=embeddings)
        upsert_duration_ms = (time.perf_counter() - upsert_started_at) * 1000

        if upserted > 0:
            self._record_summary_embedding_metadata(
                summaries=summaries,
                embeddings=embeddings,
                status="ready",
                error_message=None,
            )
        else:
            self._record_summary_embedding_metadata(
                summaries=summaries,
                embeddings=[],
                status="failed",
                error_message="vector_upsert_failed",
            )

        if upserted > 0:
            logger.info("RAG summary vector upsert: documents=%s duration_ms=%.2f", upserted, upsert_duration_ms)
        return upserted

    def _record_chunk_embedding_metadata(
        self,
        chunks: List[Dict[str, object]],
        embeddings: List[List[float]],
        status: str,
        error_message: Optional[str],
    ) -> None:
        """Сохранить технические метаданные векторной индексации чанков в БД."""
        if not chunks:
            return

        safe_status = status if status in {"ready", "failed", "stale"} else "failed"
        model_name = str(ai_settings.AI_RAG_VECTOR_EMBEDDING_MODEL or "unknown")[:128]
        vector_dim = len(embeddings[0]) if embeddings else 0
        safe_error = (error_message or "")[:255] or None

        query = """
            INSERT INTO rag_chunk_embeddings
                (document_id, chunk_index, embedding_model, embedding_dim, embedding_hash, embedding_status, error_message, embedded_at, updated_at)
            VALUES (%s, %s, %s, %s, %s, %s, %s, NOW(), NOW())
            ON DUPLICATE KEY UPDATE
                embedding_dim = VALUES(embedding_dim),
                embedding_hash = VALUES(embedding_hash),
                embedding_status = VALUES(embedding_status),
                error_message = VALUES(error_message),
                embedded_at = NOW(),
                updated_at = NOW()
        """

        params: List[Tuple[int, int, str, int, str, str, Optional[str]]] = []
        for index, chunk in enumerate(chunks):
            document_id = int(chunk.get("document_id") or 0)
            chunk_index = int(chunk.get("chunk_index") or 0)
            if document_id <= 0:
                continue

            if embeddings and index < len(embeddings):
                vector = embeddings[index]
                vector_hash = hashlib.sha256(
                    ",".join(f"{float(value):.6f}" for value in vector).encode("utf-8")
                ).hexdigest()
            else:
                vector_hash = hashlib.sha256(
                    str(chunk.get("chunk_text") or "").encode("utf-8")
                ).hexdigest()

            params.append(
                (
                    document_id,
                    chunk_index,
                    model_name,
                    vector_dim,
                    vector_hash,
                    safe_status,
                    safe_error,
                )
            )

        if not params:
            return

        try:
            for batch_start in range(0, len(params), _RAG_EMBEDDING_UPSERT_BATCH_SIZE):
                batch = params[batch_start : batch_start + _RAG_EMBEDDING_UPSERT_BATCH_SIZE]
                for attempt in range(1, _RAG_EMBEDDING_UPSERT_MAX_RETRIES + 1):
                    try:
                        with database.get_db_connection() as conn:
                            with database.get_cursor(conn) as cursor:
                                cursor.executemany(query, batch)
                        break
                    except Exception as exc:
                        mysql_errno = getattr(exc, "errno", None)
                        is_retryable = mysql_errno in _MYSQL_RETRYABLE_ERRNOS
                        if not is_retryable or attempt >= _RAG_EMBEDDING_UPSERT_MAX_RETRIES:
                            raise

                        delay = _RAG_EMBEDDING_RETRY_BASE_DELAY_SECONDS * (2 ** (attempt - 1))
                        logger.warning(
                            "Повтор сохранения rag_chunk_embeddings после временной ошибки БД: errno=%s attempt=%s/%s batch_size=%s sleep=%.2fs",
                            mysql_errno,
                            attempt,
                            _RAG_EMBEDDING_UPSERT_MAX_RETRIES,
                            len(batch),
                            delay,
                        )
                        time.sleep(delay)
        except Exception as exc:
            logger.warning("Не удалось сохранить rag_chunk_embeddings: %s", exc)

    def _record_summary_embedding_metadata(
        self,
        summaries: List[Dict[str, object]],
        embeddings: List[List[float]],
        status: str,
        error_message: Optional[str],
    ) -> None:
        """Сохранить технические метаданные векторной индексации summary документов в БД."""
        if not summaries:
            return

        safe_status = status if status in {"ready", "failed", "stale"} else "failed"
        model_name = str(ai_settings.AI_RAG_VECTOR_EMBEDDING_MODEL or "unknown")[:128]
        vector_dim = len(embeddings[0]) if embeddings else 0
        safe_error = (error_message or "")[:255] or None

        query = """
            INSERT INTO rag_summary_embeddings
                (document_id, embedding_model, embedding_dim, embedding_hash, embedding_status, error_message, embedded_at, updated_at)
            VALUES (%s, %s, %s, %s, %s, %s, NOW(), NOW())
            ON DUPLICATE KEY UPDATE
                embedding_dim = VALUES(embedding_dim),
                embedding_hash = VALUES(embedding_hash),
                embedding_status = VALUES(embedding_status),
                error_message = VALUES(error_message),
                embedded_at = NOW(),
                updated_at = NOW()
        """

        params: List[Tuple[int, str, int, str, str, Optional[str]]] = []
        for index, summary in enumerate(summaries):
            document_id = int(summary.get("document_id") or 0)
            if document_id <= 0:
                continue

            if embeddings and index < len(embeddings):
                vector = embeddings[index]
                vector_hash = hashlib.sha256(
                    ",".join(f"{float(value):.6f}" for value in vector).encode("utf-8")
                ).hexdigest()
            else:
                vector_hash = hashlib.sha256(
                    str(summary.get("summary_text") or "").encode("utf-8")
                ).hexdigest()

            params.append(
                (
                    document_id,
                    model_name,
                    vector_dim,
                    vector_hash,
                    safe_status,
                    safe_error,
                )
            )

        if not params:
            return

        try:
            for batch_start in range(0, len(params), _RAG_EMBEDDING_UPSERT_BATCH_SIZE):
                batch = params[batch_start : batch_start + _RAG_EMBEDDING_UPSERT_BATCH_SIZE]
                for attempt in range(1, _RAG_EMBEDDING_UPSERT_MAX_RETRIES + 1):
                    try:
                        with database.get_db_connection() as conn:
                            with database.get_cursor(conn) as cursor:
                                cursor.executemany(query, batch)
                        break
                    except Exception as exc:
                        mysql_errno = getattr(exc, "errno", None)
                        is_retryable = mysql_errno in _MYSQL_RETRYABLE_ERRNOS
                        if not is_retryable or attempt >= _RAG_EMBEDDING_UPSERT_MAX_RETRIES:
                            raise

                        delay = _RAG_EMBEDDING_RETRY_BASE_DELAY_SECONDS * (2 ** (attempt - 1))
                        logger.warning(
                            "Повтор сохранения rag_summary_embeddings после временной ошибки БД: errno=%s attempt=%s/%s batch_size=%s sleep=%.2fs",
                            mysql_errno,
                            attempt,
                            _RAG_EMBEDDING_UPSERT_MAX_RETRIES,
                            len(batch),
                            delay,
                        )
                        time.sleep(delay)
        except Exception as exc:
            logger.warning("Не удалось сохранить rag_summary_embeddings: %s", exc)

    def _set_vector_document_status(self, document_id: int, status: str) -> int:
        """Синхронизировать статус документа в векторном индексе."""
        if not self._is_vector_search_enabled():
            return 0

        vector_index = self._get_vector_index()
        if vector_index is None:
            return 0
        chunks_status_result = vector_index.mark_document_status(document_id=document_id, status=status)
        if ai_settings.is_rag_summary_vector_enabled():
            vector_index.mark_summary_status(document_id=document_id, status=status)
        return chunks_status_result

    def _delete_vector_document(self, document_id: int) -> int:
        """Удалить векторные точки документа при hard-delete."""
        if not self._is_vector_search_enabled():
            return 0

        vector_index = self._get_vector_index()
        if vector_index is None:
            return 0
        deleted = vector_index.delete_document_points(document_id=document_id)
        if ai_settings.is_rag_summary_vector_enabled():
            vector_index.delete_summary_points(document_id=document_id)
        return deleted

    def backfill_vector_index(
        self,
        batch_size: int = 100,
        source_type: Optional[str] = None,
        dry_run: bool = False,
        max_documents: Optional[int] = None,
        target: str = "both",
    ) -> Dict[str, int]:
        """Выполнить пакетное заполнение векторных индексов (chunks/summaries/both)."""
        stats = {
            "documents_total": 0,
            "documents_processed": 0,
            "chunks_indexed": 0,
            "summaries_indexed": 0,
            "errors": 0,
        }
        if not self._is_vector_search_enabled():
            return stats

        safe_batch_size = max(1, int(batch_size))
        normalized_target = str(target or "both").strip().lower()
        if normalized_target not in {"chunks", "summaries", "both"}:
            raise ValueError("Некорректный target для backfill: ожидается chunks|summaries|both")

        include_chunks = normalized_target in {"chunks", "both"}
        include_summaries = normalized_target in {"summaries", "both"}
        if include_summaries and not ai_settings.is_rag_summary_vector_enabled():
            include_summaries = False

        grouped_chunks = self._load_backfill_chunks(source_type=source_type) if include_chunks else {}
        grouped_summaries = self._load_backfill_summaries(source_type=source_type) if include_summaries else {}

        document_ids = sorted(set(grouped_chunks.keys()) | set(grouped_summaries.keys()))
        if max_documents is not None and int(max_documents) > 0:
            document_ids = document_ids[: int(max_documents)]

        stats["documents_total"] = len(document_ids)

        for document_id in document_ids:
            chunks = grouped_chunks.get(document_id, [])
            summary_rows = grouped_summaries.get(document_id, [])
            if not chunks and not summary_rows:
                continue

            try:
                if dry_run:
                    stats["chunks_indexed"] += len(chunks)
                    stats["summaries_indexed"] += len(summary_rows)
                else:
                    if include_chunks and chunks:
                        for start in range(0, len(chunks), safe_batch_size):
                            batch = chunks[start : start + safe_batch_size]
                            stats["chunks_indexed"] += self._upsert_vectors_for_chunks(batch)
                    if include_summaries and summary_rows:
                        for start in range(0, len(summary_rows), safe_batch_size):
                            batch = summary_rows[start : start + safe_batch_size]
                            stats["summaries_indexed"] += self._upsert_vectors_for_summaries(batch)
                stats["documents_processed"] += 1
            except Exception as exc:
                stats["errors"] += 1
                logger.warning("Ошибка backfill vector index для document_id=%s: %s", document_id, exc)

        return stats

    @staticmethod
    def _load_backfill_chunks(source_type: Optional[str]) -> Dict[int, List[Dict[str, object]]]:
        """Загрузить активные чанки документов для chunk backfill."""
        query = """
            SELECT d.id, d.filename, d.source_type, c.chunk_index, c.chunk_text
            FROM rag_documents d
            JOIN rag_chunks c ON c.document_id = d.id
            WHERE d.status = 'active'
        """
        params: List[object] = []
        if source_type:
            query += " AND d.source_type = %s"
            params.append(source_type)
        query += " ORDER BY d.id ASC, c.chunk_index ASC"

        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                cursor.execute(query, tuple(params))
                rows = cursor.fetchall() or []

        grouped: Dict[int, List[Dict[str, object]]] = {}
        for row in rows:
            document_id = int(row.get("id") or 0)
            if document_id <= 0:
                continue
            grouped.setdefault(document_id, []).append(
                {
                    "document_id": document_id,
                    "chunk_index": int(row.get("chunk_index") or 0),
                    "filename": str(row.get("filename") or "document"),
                    "chunk_text": str(row.get("chunk_text") or ""),
                    "status": "active",
                }
            )
        return grouped

    @staticmethod
    def _load_backfill_summaries(source_type: Optional[str]) -> Dict[int, List[Dict[str, object]]]:
        """Загрузить активные summary документов для summary backfill."""
        query = """
            SELECT d.id, d.filename, d.source_type, s.summary_text
            FROM rag_documents d
            JOIN rag_document_summaries s ON s.document_id = d.id
            WHERE d.status = 'active'
        """
        params: List[object] = []
        if source_type:
            query += " AND d.source_type = %s"
            params.append(source_type)
        query += " ORDER BY d.id ASC"

        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                cursor.execute(query, tuple(params))
                rows = cursor.fetchall() or []

        grouped: Dict[int, List[Dict[str, object]]] = {}
        for row in rows:
            document_id = int(row.get("id") or 0)
            summary_text = str(row.get("summary_text") or "").strip()
            if document_id <= 0 or not summary_text:
                continue

            grouped.setdefault(document_id, []).append(
                {
                    "document_id": document_id,
                    "filename": str(row.get("filename") or "document"),
                    "summary_text": summary_text,
                    "status": "active",
                }
            )

        return grouped

    def _tokenize(self, text: str) -> List[str]:
        """Токенизировать текст для lexical retrieval c опциональной нормализацией RU."""
        tokens = _TOKEN_RE.findall((text or "").lower())
        if not tokens:
            return []

        if not ai_settings.is_rag_ru_normalization_enabled():
            return tokens

        return [self._normalize_token(token) for token in tokens if token]

    def _normalize_token(self, token: str) -> str:
        """Нормализовать токен с кэшем (лемматизация/стемминг для русского)."""
        safe_token = (token or "").strip().lower()
        if not safe_token:
            return ""

        cached = self._normalized_token_cache.get(safe_token)
        if cached is not None:
            return cached

        if not _CYRILLIC_TOKEN_RE.search(safe_token):
            self._normalized_token_cache[safe_token] = safe_token
            return safe_token

        mode = ai_settings.get_rag_ru_normalization_mode()
        normalized = safe_token
        if mode in {"lemma_then_stem", "lemma_only"}:
            normalized = self._lemmatize_ru_token(normalized)
        if mode in {"lemma_then_stem", "stem_only"}:
            normalized = self._stem_ru_token(normalized)

        normalized = normalized or safe_token
        self._normalized_token_cache[safe_token] = normalized
        return normalized

    def _lemmatize_ru_token(self, token: str) -> str:
        """Лемматизировать русский токен через pymorphy3 с graceful fallback."""
        analyzer = self._get_ru_morph_analyzer()
        if analyzer is None:
            return token

        try:
            parsed = analyzer.parse(token)
            if parsed:
                normal_form = str(getattr(parsed[0], "normal_form", "") or "").strip().lower()
                if normal_form:
                    return normal_form
        except Exception:
            return token

        return token

    def _stem_ru_token(self, token: str) -> str:
        """Применить стемминг русского токена через snowballstemmer."""
        stemmer = self._get_ru_stemmer()
        if stemmer is None:
            return token

        try:
            stemmed = stemmer.stemWord(token)
            return str(stemmed or token).strip().lower()
        except Exception:
            return token

    def _get_ru_morph_analyzer(self) -> Optional[object]:
        """Ленивая инициализация pymorphy3 MorphAnalyzer."""
        if self._ru_morph_analyzer is not None:
            return self._ru_morph_analyzer

        try:
            import pymorphy3

            self._ru_morph_analyzer = pymorphy3.MorphAnalyzer()
            return self._ru_morph_analyzer
        except Exception as exc:
            if not self._normalization_dependency_warning_logged:
                logger.warning("RU-нормализация: pymorphy3 недоступен, fallback без лемматизации: %s", exc)
                self._normalization_dependency_warning_logged = True
            return None

    def _get_ru_stemmer(self) -> Optional[object]:
        """Ленивая инициализация snowballstemmer для русского."""
        if self._ru_stemmer is not None:
            return self._ru_stemmer

        try:
            import snowballstemmer

            self._ru_stemmer = snowballstemmer.stemmer("russian")
            return self._ru_stemmer
        except Exception as exc:
            if not self._normalization_dependency_warning_logged:
                logger.warning("RU-нормализация: snowballstemmer недоступен, fallback без стемминга: %s", exc)
                self._normalization_dependency_warning_logged = True
            return None

    @staticmethod
    def _score_corpus_bm25(corpus_tokens: List[List[str]], query_tokens: List[str]) -> List[float]:
        """Вычислить BM25-score для корпуса документов по токенам запроса."""
        if not corpus_tokens or not query_tokens:
            return [0.0 for _ in corpus_tokens]

        safe_corpus = [tokens if tokens else [""] for tokens in corpus_tokens]

        if BM25Okapi is not None:
            try:
                bm25 = BM25Okapi(
                    safe_corpus,
                    k1=max(0.01, float(ai_settings.AI_RAG_BM25_K1)),
                    b=max(0.0, min(1.0, float(ai_settings.AI_RAG_BM25_B))),
                )
                return [max(0.0, float(score)) for score in bm25.get_scores(query_tokens)]
            except Exception:
                pass

        k1 = max(0.01, float(ai_settings.AI_RAG_BM25_K1))
        b = max(0.0, min(1.0, float(ai_settings.AI_RAG_BM25_B)))
        doc_count = len(safe_corpus)
        avg_doc_len = sum(len(doc_tokens) for doc_tokens in safe_corpus) / max(doc_count, 1)

        doc_freq: Dict[str, int] = {}
        term_freq_by_doc: List[Counter[str]] = []
        for doc_tokens in safe_corpus:
            tf = Counter(doc_tokens)
            term_freq_by_doc.append(tf)
            for token in tf.keys():
                doc_freq[token] = doc_freq.get(token, 0) + 1

        query_tf = Counter(query_tokens)
        scores: List[float] = [0.0 for _ in safe_corpus]

        for token, qf in query_tf.items():
            df = doc_freq.get(token, 0)
            if df <= 0:
                continue

            idf = math.log(1.0 + ((doc_count - df + 0.5) / (df + 0.5)))
            for idx, tf in enumerate(term_freq_by_doc):
                freq = tf.get(token, 0)
                if freq <= 0:
                    continue
                doc_len = len(safe_corpus[idx])
                denominator = freq + (k1 * (1.0 - b + (b * doc_len / max(avg_doc_len, 1e-9))))
                if denominator <= 0:
                    continue
                scores[idx] += idf * ((freq * (k1 + 1.0)) / denominator) * qf

        return [max(0.0, float(score)) for score in scores]

    @staticmethod
    def _score_chunk(chunk_text: str, question_tokens: List[str]) -> float:
        """Оценить релевантность чанка по вхождению токенов запроса."""
        low_chunk = (chunk_text or "").lower()
        if not low_chunk:
            return 0.0

        unique_tokens = set(question_tokens)
        matches = sum(1 for token in unique_tokens if token in low_chunk)
        if matches == 0:
            return 0.0

        token_coverage = matches / max(len(unique_tokens), 1)
        density_bonus = min(matches / 10.0, 0.3)
        return token_coverage + density_bonus

    @staticmethod
    def _get_corpus_version() -> int:
        """Получить текущую версию корпуса базы знаний."""
        with database.get_db_connection() as conn:
            with database.get_cursor(conn) as cursor:
                cursor.execute("SELECT COALESCE(MAX(id), 0) AS version_id FROM rag_corpus_version")
                row = cursor.fetchone() or {"version_id": 0}
                return int(row.get("version_id", 0))

    @staticmethod
    def _bump_corpus_version(cursor, reason: str) -> None:
        """Увеличить версию корпуса для инвалидации кэша и retrieval-состояния."""
        cursor.execute(
            """
            INSERT INTO rag_corpus_version (reason, created_at)
            VALUES (%s, NOW())
            """,
            (reason[:255],),
        )

    def _clear_expired_cache(self) -> None:
        """Очистить протухшие элементы кэша."""
        now = time.time()
        expired = [key for key, val in self._answer_cache.items() if val.expires_at <= now]
        for key in expired:
            self._answer_cache.pop(key, None)

    @staticmethod
    def _log_query(user_id: int, query: str, cache_hit: bool, chunks_count: int) -> None:
        """Записать факт запроса к базе знаний в БД."""
        try:
            with database.get_db_connection() as conn:
                with database.get_cursor(conn) as cursor:
                    cursor.execute(
                        """
                        INSERT INTO rag_query_log (user_id, query_text, cache_hit, chunks_count, created_at)
                        VALUES (%s, %s, %s, %s, NOW())
                        """,
                        (user_id, query[:1000], 1 if cache_hit else 0, chunks_count),
                    )
        except Exception as exc:
            logger.warning("Не удалось записать rag_query_log: %s", exc)


_rag_service_instance: Optional[RagKnowledgeService] = None


def get_rag_service() -> RagKnowledgeService:
    """Получить singleton-экземпляр RAG-сервиса."""
    global _rag_service_instance
    if _rag_service_instance is None:
        _rag_service_instance = RagKnowledgeService()
    return _rag_service_instance


def preload_rag_runtime_dependencies() -> Dict[str, bool]:
    """Прогреть lazy-зависимости RAG на старте процесса бота."""
    started_at = time.perf_counter()
    logger.info("RAG preload: start")

    preload_result: Dict[str, bool] = {
        "vector_provider_ready": False,
        "vector_index_ready": False,
        "ru_morph_ready": False,
        "ru_stemmer_ready": False,
    }

    vector_enabled = False
    ru_normalization_enabled = bool(ai_settings.is_rag_ru_normalization_enabled())
    ru_normalization_mode = ai_settings.get_rag_ru_normalization_mode()
    status = "ok"

    try:
        rag_service = get_rag_service()

        vector_enabled = bool(rag_service._is_vector_search_enabled())
        if vector_enabled:
            preload_result["vector_provider_ready"] = rag_service._get_embedding_provider() is not None
            preload_result["vector_index_ready"] = rag_service._get_vector_index() is not None

        if ru_normalization_enabled and ru_normalization_mode in {"lemma_then_stem", "lemma_only"}:
            preload_result["ru_morph_ready"] = rag_service._get_ru_morph_analyzer() is not None

        if ru_normalization_enabled and ru_normalization_mode in {"lemma_then_stem", "stem_only"}:
            preload_result["ru_stemmer_ready"] = rag_service._get_ru_stemmer() is not None
    except Exception:
        status = "failed"
        logger.exception("RAG preload: failed")
    finally:
        duration_ms = int((time.perf_counter() - started_at) * 1000)
        logger.info(
            "RAG preload: done status=%s duration_ms=%s vector_enabled=%s vector_provider_ready=%s vector_index_ready=%s ru_normalization_enabled=%s ru_normalization_mode=%s ru_morph_ready=%s ru_stemmer_ready=%s",
            status,
            duration_ms,
            vector_enabled,
            preload_result["vector_provider_ready"],
            preload_result["vector_index_ready"],
            ru_normalization_enabled,
            ru_normalization_mode,
            preload_result["ru_morph_ready"],
            preload_result["ru_stemmer_ready"],
        )

    return preload_result


def reset_rag_service() -> None:
    """Сбросить singleton-экземпляр RAG-сервиса (для тестов)."""
    global _rag_service_instance
    _rag_service_instance = None
